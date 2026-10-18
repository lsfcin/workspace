# T1 docx template: a filled form reads back as the fields written, with the template's defaults gone and its table found even inside a content control.
import pathlib, sys

import pytest

docx = pytest.importorskip("docx")
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "docs"))
import docx_template
from docx.oxml.ns import qn

LABELS = [("NOME:", "nome", "inline"), ("Teórica:", "carga", "cell"), ("EMENTA:", "ementa", "block")]


def _template(path, in_sdt=False):
    doc = docx.Document()
    table = doc.add_table(rows=3, cols=1)
    table.cell(0, 0).text = "NOME: NÃO TEM"
    table.cell(1, 0).text = "Teórica: 60h"
    table.cell(1, 0).add_paragraph("Prática:")
    table.cell(2, 0).text = "EMENTA:"
    if in_sdt:
        tbl = table._tbl
        sdt = tbl.makeelement(qn("w:sdt"), {})
        content = sdt.makeelement(qn("w:sdtContent"), {})
        tbl.addprevious(sdt)
        sdt.append(content)
        content.append(tbl)
    doc.save(path)
    return str(path)


@pytest.mark.parametrize("in_sdt", [False, True])
def test_a_filled_form_reads_back_as_the_fields_written(tmp_path, in_sdt):
    fields = {"nome": "Arquiteturas", "carga": " 45h\nPrática: 15h", "ementa": "linha um\nlinha dois"}
    out = str(tmp_path / "out.docx")
    assert docx_template.fill(_template(tmp_path / "t.docx", in_sdt), LABELS, fields, out) == []
    back = docx_template.read(out, LABELS)
    assert back == {"nome": "Arquiteturas", "carga": "45h\nPrática: 15h", "ementa": "linha um\nlinha dois"}


def test_an_inline_value_replaces_the_template_default_instead_of_following_it(tmp_path):
    out = str(tmp_path / "out.docx")
    docx_template.fill(_template(tmp_path / "t.docx"), LABELS, {"nome": "X"}, out)
    assert "NÃO TEM" not in docx.Document(out).tables[0].cell(0, 0).text


def test_a_field_whose_label_has_no_cell_is_reported_not_dropped(tmp_path):
    labels = LABELS + [("OBJETIVOS:", "objetivos", "block")]
    out = str(tmp_path / "out.docx")
    assert docx_template.fill(_template(tmp_path / "t.docx"), labels, {"objetivos": "y"}, out) == ["objetivos"]


def test_a_forced_font_reaches_every_run_of_the_table(tmp_path):
    out = str(tmp_path / "out.docx")
    docx_template.fill(_template(tmp_path / "t.docx"), LABELS, {"nome": "X"}, out, font="Times New Roman")
    runs = list(docx.Document(out).tables[0]._tbl.iter(qn("w:r")))
    assert runs and all(r.find(qn("w:rPr")).find(qn("w:rFonts")).get(qn("w:ascii")) == "Times New Roman" for r in runs)


def test_a_numbered_list_reads_with_the_numbers_the_reader_sees(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("CONTEÚDOS")
    for text in ("um", "dois"):
        doc.add_paragraph(text, style="List Number")
    path = str(tmp_path / "l.docx")
    doc.save(path)
    assert docx_template.paragraphs(path) == ["CONTEÚDOS", "1. um", "2. dois"]
