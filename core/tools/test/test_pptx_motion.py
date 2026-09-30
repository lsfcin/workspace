# T1 pptx motion: a build shows what it means to show, lint names what a room would see wrong, and flatten turns every click into a page.
import io, pathlib, sys

import pytest

pptx = pytest.importorskip("pptx")
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "slides"))
import pptx_check, pptx_motion
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches


def _deck():
    prs = pptx.Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def _node(slide, label, x):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(3), Inches(2), Inches(2))
    sh.text_frame.text = label
    return sh


def _reopen(prs):
    buf = io.BytesIO(); prs.save(buf); buf.seek(0)
    return pptx.Presentation(buf)


def test_an_entrance_animates_the_whole_shape_not_only_its_label():
    prs = _deck(); s = prs.slides.add_slide(prs.slide_layouts[6])
    a = _node(s, "RNN", 1)
    pptx_motion.timeline(s, [[pptx_motion.fade_in(a)]])
    assert pptx_check.lint(_reopen(prs)) == []
    s._element.find(".//" + pptx_motion.q(pptx_motion.P, "bldP")).attrib.pop("animBg")
    assert any("animBg" in f for f in pptx_check.lint(prs))


def test_timing_ids_are_unique_across_clicks_and_triggers():
    prs = _deck(); s = prs.slides.add_slide(prs.slide_layouts[6])
    a, b = _node(s, "a", 1), _node(s, "b", 5)
    pptx_motion.timeline(s, [[pptx_motion.fade_in(a), pptx_motion.dim(b)], [pptx_motion.grow(a, 150)]],
                         triggers=[(b, [pptx_motion.spin(b, 90)])])
    ids = [c.get("id") for c in s._element.iter(pptx_motion.q(pptx_motion.P, "cTn"))]
    assert len(ids) == len(set(ids))


def test_a_move_and_a_scale_on_one_click_with_different_durations_is_flagged():
    prs = _deck(); s = prs.slides.add_slide(prs.slide_layouts[6])
    a = _node(s, "LSTM", 1)
    pptx_motion.timeline(s, [[pptx_motion.move(a, "M 0 0 L 0.2 0 E", dur=800), pptx_motion.grow(a, 150, dur=600)]])
    assert any("desync" in f for f in pptx_check.lint(prs))


def test_a_move_over_a_third_of_the_slide_is_flagged():
    prs = _deck(); s = prs.slides.add_slide(prs.slide_layouts[6])
    a = _node(s, "h", 1)
    pptx_motion.timeline(s, [[pptx_motion.move(a, "M 0 0 L 0.6 0 E")]])
    assert any("third" in f for f in pptx_check.lint(prs))


def test_a_morph_with_nothing_paired_is_flagged_and_a_named_pair_is_not():
    prs = _deck()
    s1, s2 = prs.slides.add_slide(prs.slide_layouts[6]), prs.slides.add_slide(prs.slide_layouts[6])
    pptx_motion.morph_key(_node(s1, "x", 1), "x")
    b = _node(s2, "x", 6)
    pptx_motion.transition(s2, "morph")
    assert any("nothing is paired" in f for f in pptx_check.lint(prs))
    pptx_motion.morph_key(b, "x")
    assert pptx_check.lint(_reopen(prs)) == []


def test_flatten_gives_one_page_per_click_in_order_with_each_state():
    prs = _deck()
    first = prs.slides.add_slide(prs.slide_layouts[6])
    s = prs.slides.add_slide(prs.slide_layouts[6])
    last = prs.slides.add_slide(prs.slide_layouts[6])
    a, b = _node(s, "a", 1), _node(s, "b", 5)
    x0 = b.left
    pptx_motion.timeline(s, [[pptx_motion.fade_in(a)], [pptx_motion.move(b, "M 0 0 L 0.1 0 E")], [pptx_motion.fade_out(a)]])
    assert pptx_check.flatten(prs) == 3
    prs = _reopen(prs)
    names = [[sh.text_frame.text for sh in sl.shapes] for sl in prs.slides]
    assert names == [[], ["b"], ["a", "b"], ["a", "b"], ["b"], []]
    moved = [sh for sh in prs.slides[3].shapes if sh.text_frame.text == "b"][0]
    assert moved.left == x0 + int(0.1 * prs.slide_width)
    assert all(sl._element.find(pptx_motion.q(pptx_motion.P, "timing")) is None for sl in prs.slides)


def test_a_back_link_is_a_plain_jump_to_the_named_slide():
    prs = _deck()
    menu, topic = prs.slides.add_slide(prs.slide_layouts[6]), prs.slides.add_slide(prs.slide_layouts[6])
    btn = pptx_motion.link(_node(topic, "voltar", 1), menu)
    assert btn.click_action.target_slide == menu
    h = btn._element.find(".//" + pptx_motion.q(pptx_motion.A, "hlinkClick"))
    assert h.get("action") == "ppaction://hlinksldjump"


def test_the_first_group_can_play_on_its_own_once_the_slide_arrives():
    prs = _deck(); s = prs.slides.add_slide(prs.slide_layouts[6])
    a, b = _node(s, "a", 1), _node(s, "b", 5)
    pptx_motion.timeline(s, [[pptx_motion.fade_in(a)], [pptx_motion.fade_in(b)]], auto_first=True)
    xml = _xml(s)
    assert xml.count('evt="onBegin"') == 1 and xml.count('nodeType="afterEffect"') == 1 and xml.count('nodeType="clickEffect"') == 1


def test_text_at_alpha_zero_is_exempt_from_the_size_floor():
    prs = _deck(); s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    r = tb.text_frame.paragraphs[0].add_run(); r.text = "memória"; r.font.size = pptx.util.Pt(8)
    r.font.color.rgb = pptx.dml.color.RGBColor(0, 0, 0)
    assert any("below 14" in f for f in pptx_check.lint(prs))
    clr = r._r.find(".//" + pptx_motion.q(pptx_motion.A, "srgbClr"))
    clr.append(clr.makeelement(pptx_motion.q(pptx_motion.A, "alpha"), {"val": "0"}))
    assert pptx_check.lint(prs) == []


def test_a_slide_over_the_item_budget_is_flagged_and_one_at_it_is_not():
    prs = _deck(); s = prs.slides.add_slide(prs.slide_layouts[6])
    for k in range(pptx_check.MAX_ITEMS):
        s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(k % 15 * 0.8), Inches(k // 15 * 0.7), Inches(0.5), Inches(0.5))
    assert not any("items" in f for f in pptx_check.lint(prs))
    grp = s.shapes.add_group_shape()   # a group's children count: the online editor counts them too
    grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(0.5), Inches(0.5))
    assert any(f"{pptx_check.MAX_ITEMS + 1} items" in f for f in pptx_check.lint(prs))


def _xml(slide):
    from lxml import etree
    return etree.tostring(slide._element).decode()
