# pptx_masters.py — a .pptx with one master per section colour, each carrying the same layouts
import copy

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.parts.slide import SlideLayoutPart, SlideMasterPart
from pptx.util import Emu

import pptx_layouts, pptx_shapes as S, pptx_theme

# Why a .pptx at all: the Slides API cannot create or duplicate a layout or a master
# (core/tools/slides/SPECS.md), and Drive's pptx->Slides conversion keeps both, with shapes
# still pointing at THEME slots — so a section is a master, and its colour lives in one place.
P = {"p": S.P, "r": S.R}
FIRST_ID = 2147483648  # master and layout ids share one space and must start here


def _clear_layouts(master_part) -> None:
    lst = master_part._element.find("p:sldLayoutIdLst", P)
    for e in list(lst):
        master_part.drop_rel(e.get(f"{{{S.R}}}id"))
        lst.remove(e)


def _append_id(lst, next_id: int, rid: str, tag: str) -> None:
    e = etree.SubElement(lst, f"{{{S.P}}}{tag}")
    e.set("id", str(next_id))
    e.set(f"{{{S.R}}}id", rid)


def build(spec: dict):
    """The Presentation for a spec: len(sections) masters x len(layouts) layouts, plus one
    sample slide per layout (first section) and one section slide per colour, so it previews."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(S.W), Emu(S.H)
    pkg, first = prs.part.package, prs.slide_masters[0].part
    stock_theme = first.part_related_by(RT.THEME)
    _clear_layouts(first)
    blank_master = copy.deepcopy(first._element)
    ids = iter(range(FIRST_ID, FIRST_ID + 100000))
    masters = prs.part._element.find("p:sldMasterIdLst", P)
    masters.find("p:sldMasterId", P).set("id", str(next(ids)))
    layouts = pptx_layouts.build(spec)
    n = 0
    for k, sec in enumerate(spec["sections"]):
        blob = pptx_theme.rewrite(stock_theme.blob, sec["name"], pptx_theme.slots(spec, sec), spec["fonts"])
        if k == 0:
            mp = first
            stock_theme._blob = blob
        else:
            mp = SlideMasterPart(PackURI(f"/ppt/slideMasters/slideMaster{k + 1}.xml"), CT.PML_SLIDE_MASTER, pkg,
                                 copy.deepcopy(blank_master))
            mp.relate_to(Part(PackURI(f"/ppt/theme/theme{k + 1}.xml"), CT.OFC_THEME, pkg, blob), RT.THEME)
            _append_id(masters, next(ids), prs.part.relate_to(mp, RT.SLIDE_MASTER), "sldMasterId")
        mp._element.find("p:cSld", P).set("name", sec["name"])
        lst = mp._element.find("p:sldLayoutIdLst", P)
        for _, xml in layouts:
            n += 1
            lp = SlideLayoutPart(PackURI(f"/ppt/slideLayouts/slideLayout{n}.xml"), CT.PML_SLIDE_LAYOUT, pkg, parse_xml(xml))
            lp.relate_to(mp, RT.SLIDE_MASTER)
            _append_id(lst, next(ids), mp.relate_to(lp, RT.SLIDE_LAYOUT), "sldLayoutId")
    _samples(prs, spec)
    return prs


def _sample(prs, layout) -> None:
    """A slide on `layout` whose text slots say what goes there, so a preview is not blank."""
    slide = prs.slides.add_slide(layout)
    said = {ph.placeholder_format.idx: ph.text_frame.text for ph in layout.placeholders if ph.has_text_frame}
    for ph in slide.placeholders:
        if ph.has_text_frame and said.get(ph.placeholder_format.idx):
            ph.text_frame.text = said[ph.placeholder_format.idx]


def _samples(prs, spec: dict) -> None:
    section_kind = next((j for j, p in enumerate(spec["layouts"]) if p["kind"] == "section"), None)
    for layout in prs.slide_masters[0].slide_layouts:
        _sample(prs, layout)
    if section_kind is not None:
        for m in list(prs.slide_masters)[1:]:
            _sample(prs, m.slide_layouts[section_kind])


def summary(prs) -> str:
    ms = list(prs.slide_masters)
    return f"{len(ms)} masters x {len(ms[0].slide_layouts)} layouts, {len(prs.slides)} sample slides"
